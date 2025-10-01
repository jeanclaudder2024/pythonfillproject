#!/usr/bin/env python3
"""
Debug Mismatch Script
Checks for mismatches between placeholder names and replacement data keys
"""

import os
from document_processor import DocumentProcessor
from docx import Document

def debug_mismatch():
    """Debug mismatches between placeholders and replacement data"""
    print("🔍 DEBUGGING PLACEHOLDER-DATA MISMATCH")
    print("=" * 60)
    
    # Use one of the uploaded documents
    test_file = "uploads/000d8889-db1e-4707-bc66-a95f78933dc0_Proof_of_Product_Template_with_placeholders.docx"
    
    if not os.path.exists(test_file):
        print(f"❌ Test file not found: {test_file}")
        return
    
    print(f"📄 Testing with: {os.path.basename(test_file)}")
    
    # Initialize processor
    processor = DocumentProcessor()
    
    # Load document
    doc = Document(test_file)
    
    # Find placeholders
    placeholders = processor.find_placeholders(doc)
    print(f"\n🔍 Found {len(placeholders)} placeholders")
    
    # Generate replacement data
    replacement_data = processor.generate_random_replacement_data(placeholders)
    print(f"📊 Generated {len(replacement_data)} replacement data keys")
    
    # Check for matches
    print("\n🎯 CHECKING MATCHES...")
    matched = 0
    unmatched_placeholders = []
    
    for placeholder in placeholders:
        if placeholder in replacement_data:
            matched += 1
        else:
            unmatched_placeholders.append(placeholder)
    
    print(f"✅ Matched: {matched}/{len(placeholders)} placeholders")
    print(f"❌ Unmatched: {len(unmatched_placeholders)} placeholders")
    
    if unmatched_placeholders:
        print("\n🚫 UNMATCHED PLACEHOLDERS:")
        for i, placeholder in enumerate(unmatched_placeholders[:20]):  # Show first 20
            print(f"  {i+1}. {placeholder}")
        if len(unmatched_placeholders) > 20:
            print(f"  ... and {len(unmatched_placeholders) - 20} more")
    
    # Check for extra keys in replacement data
    extra_keys = [key for key in replacement_data.keys() if key not in placeholders]
    print(f"\n📊 Extra keys in replacement data: {len(extra_keys)}")
    if extra_keys:
        print("🔑 EXTRA KEYS (first 20):")
        for i, key in enumerate(extra_keys[:20]):
            print(f"  {i+1}. {key}: {str(replacement_data[key])[:50]}{'...' if len(str(replacement_data[key])) > 50 else ''}")
        if len(extra_keys) > 20:
            print(f"  ... and {len(extra_keys) - 20} more")
    
    # Test specific placeholder replacement
    print("\n🧪 TESTING SPECIFIC REPLACEMENTS...")
    
    # Get some sample text from the document
    sample_paragraphs = []
    for i, paragraph in enumerate(doc.paragraphs[:5]):
        if paragraph.text.strip():
            sample_paragraphs.append((i, paragraph.text))
    
    for para_num, para_text in sample_paragraphs:
        print(f"\nParagraph {para_num}: {para_text[:100]}...")
        
        # Check if this paragraph contains any placeholders
        import re
        found_placeholders = re.findall(r'\{([^}]+)\}', para_text)
        if found_placeholders:
            print(f"  Contains placeholders: {found_placeholders}")
            
            # Test replacement
            new_text = para_text
            for placeholder in found_placeholders:
                if placeholder in replacement_data:
                    value = str(replacement_data[placeholder])
                    new_text = new_text.replace(f"{{{placeholder}}}", value)
                    print(f"    {placeholder} -> {value[:30]}{'...' if len(value) > 30 else ''}")
                else:
                    print(f"    {placeholder} -> NO DATA AVAILABLE")
            
            if new_text != para_text:
                print(f"  After replacement: {new_text[:100]}...")
            else:
                print(f"  No changes made")
        else:
            print(f"  No placeholders found")

if __name__ == "__main__":
    debug_mismatch()