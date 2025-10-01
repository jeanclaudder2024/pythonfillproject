#!/usr/bin/env python3
"""
Find Unmapped Placeholders Script
Identifies placeholders that are not properly mapped in the random data generator
"""

import os
from document_processor import DocumentProcessor
from random_data_generator import RandomDataGenerator
from docx import Document

def find_unmapped_placeholders():
    """Find placeholders that are not properly mapped"""
    print("🔍 FINDING UNMAPPED PLACEHOLDERS")
    print("=" * 60)
    
    # Use one of the uploaded documents
    test_file = "uploads/000d8889-db1e-4707-bc66-a95f78933dc0_Proof_of_Product_Template_with_placeholders.docx"
    
    if not os.path.exists(test_file):
        print(f"❌ Test file not found: {test_file}")
        return
    
    print(f"📄 Testing with: {os.path.basename(test_file)}")
    
    # Initialize processor and generator
    processor = DocumentProcessor()
    generator = RandomDataGenerator()
    
    # Load document and find placeholders
    doc = Document(test_file)
    placeholders = processor.find_placeholders(doc)
    print(f"\n🔍 Found {len(placeholders)} placeholders")
    
    # Test each placeholder
    unmapped_placeholders = []
    fallback_placeholders = []
    
    for placeholder in placeholders:
        value = generator.get_random_value(placeholder)
        
        # Check if it's a fallback value (contains square brackets)
        if value.startswith('[') and value.endswith(']'):
            fallback_placeholders.append((placeholder, value))
        
        # Check if it's properly mapped
        normalized_key = generator._normalize_key(placeholder)
        if normalized_key not in generator.data_pools:
            unmapped_placeholders.append((placeholder, normalized_key, value))
    
    print(f"\n📊 RESULTS:")
    print(f"✅ Properly mapped: {len(placeholders) - len(unmapped_placeholders)}")
    print(f"❌ Unmapped: {len(unmapped_placeholders)}")
    print(f"🔄 Using fallback: {len(fallback_placeholders)}")
    
    if unmapped_placeholders:
        print(f"\n🚫 UNMAPPED PLACEHOLDERS:")
        for i, (original, normalized, value) in enumerate(unmapped_placeholders):
            print(f"  {i+1}. '{original}' -> '{normalized}' -> '{value}'")
    
    if fallback_placeholders:
        print(f"\n🔄 FALLBACK PLACEHOLDERS:")
        for i, (placeholder, value) in enumerate(fallback_placeholders):
            print(f"  {i+1}. '{placeholder}' -> '{value}'")
    
    # Suggest mappings for unmapped placeholders
    if unmapped_placeholders:
        print(f"\n💡 SUGGESTED MAPPINGS:")
        suggestions = {}
        
        for original, normalized, value in unmapped_placeholders:
            # Analyze the placeholder name to suggest appropriate mapping
            key = original.lower()
            
            if 'reference' in key or 'ref' in key:
                suggestions[original] = "Generate reference number (e.g., REF-2025-001)"
            elif 'certificate' in key or 'cert' in key:
                suggestions[original] = "Generate certificate number (e.g., CERT-2025-001)"
            elif 'batch' in key:
                suggestions[original] = "Generate batch number (e.g., BATCH-2025-001)"
            elif 'lot' in key:
                suggestions[original] = "Generate lot number (e.g., LOT-2025-001)"
            elif 'serial' in key:
                suggestions[original] = "Generate serial number (e.g., SN-2025-001)"
            elif 'id' in key:
                suggestions[original] = "Generate ID number (e.g., ID-2025-001)"
            elif 'code' in key:
                suggestions[original] = "Generate code (e.g., CODE-2025-001)"
            elif 'grade' in key:
                suggestions[original] = "Generate grade (e.g., Grade A, Premium)"
            elif 'specification' in key or 'spec' in key:
                suggestions[original] = "Generate specification (e.g., ISO 8217:2017)"
            elif 'standard' in key:
                suggestions[original] = "Generate standard (e.g., ASTM D975)"
            else:
                suggestions[original] = f"Add mapping for '{normalized}'"
        
        for placeholder, suggestion in suggestions.items():
            print(f"  • {placeholder}: {suggestion}")

if __name__ == "__main__":
    find_unmapped_placeholders()